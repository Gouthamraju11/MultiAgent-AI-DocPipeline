from __future__ import annotations

import csv
import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".pdf", ".docx"}


class DocumentReadError(ValueError):
    pass


def extract_text(content: bytes, filename: str) -> str:
    extension = Path(filename or "document.txt").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentReadError(
            f"Unsupported file type '{extension or 'unknown'}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if not content:
        raise DocumentReadError("The uploaded file is empty")

    try:
        if extension == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise DocumentReadError("Encrypted PDFs are not supported")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                raise DocumentReadError(
                    "No selectable text was found in the PDF. Scanned PDFs require OCR before upload."
                )
            return text

        if extension == ".docx":
            import docx

            document = docx.Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                    if any(cell.text.strip() for cell in row.cells)
                )
            return "\n".join(paragraphs)

        decoded = content.decode("utf-8-sig")
        if extension == ".csv":
            rows = csv.reader(io.StringIO(decoded))
            return "\n".join(", ".join(cell.strip() for cell in row) for row in rows)
        return decoded
    except DocumentReadError:
        raise
    except UnicodeDecodeError as exc:
        raise DocumentReadError("Text files must be UTF-8 encoded") from exc
    except Exception as exc:
        raise DocumentReadError(f"Could not read {extension} document: {exc}") from exc
