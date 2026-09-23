from __future__ import annotations

import csv
import io
import json
from pathlib import PurePosixPath

from shared.config import MAX_DOCUMENT_BYTES

SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".pdf", ".docx"}


class DocumentError(ValueError):
    pass


def extract_text(content: bytes, filename: str) -> str:
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(f"Unsupported document extension: {extension or '(none)'}")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentError("Document exceeds the 10 MB processing limit")

    if extension in {".txt", ".md"}:
        return _decode(content)
    if extension == ".json":
        try:
            return json.dumps(json.loads(_decode(content)), indent=2, ensure_ascii=False)
        except json.JSONDecodeError as exc:
            raise DocumentError(f"Invalid JSON document: {exc.msg}") from exc
    if extension == ".csv":
        source = io.StringIO(_decode(content))
        return "\n".join(" | ".join(row) for row in csv.reader(source))
    if extension == ".pdf":
        return _extract_pdf(content)
    return _extract_docx(content)


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentError("Text documents must use UTF-8 encoding") from exc


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentError("Encrypted PDFs are not supported")
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The PDF is invalid or unreadable") from exc
    if not text:
        raise DocumentError("No selectable text found; scanned PDFs require OCR")
    return text


def _extract_docx(content: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
        text = "\n".join([*paragraphs, *cells]).strip()
    except Exception as exc:
        raise DocumentError("The DOCX file is invalid or unreadable") from exc
    if not text:
        raise DocumentError("No readable text found in the DOCX file")
    return text
