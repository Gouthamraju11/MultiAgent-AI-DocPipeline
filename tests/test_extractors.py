import io
import zipfile

import pytest
from app.extractors import DocumentReadError, extract_text


def test_plain_text_requires_utf8() -> None:
    with pytest.raises(DocumentReadError, match="UTF-8"):
        extract_text(b"\xff\xfe\xfd", "notes.txt")


def test_unsupported_extension_has_clear_error() -> None:
    with pytest.raises(DocumentReadError, match="Unsupported file type"):
        extract_text(b"content", "document.pages")


def test_docx_includes_table_content() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Invoice summary")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Total"
    table.cell(0, 1).text = "$25.00"
    output = io.BytesIO()
    document.save(output)
    text = extract_text(output.getvalue(), "invoice.docx")
    assert "Invoice summary" in text
    assert "Total | $25.00" in text


def test_invalid_docx_does_not_fall_back_to_binary_decode() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("not-a-document.txt", "bad")
    with pytest.raises(DocumentReadError, match="Could not read"):
        extract_text(buffer.getvalue(), "broken.docx")
